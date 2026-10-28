"""Completa las columnas «Estado» y «Resultado obtenido» del documento 14 con lo que dejaron las pruebas.

Uso (desde la carpeta del proyecto):  python scripts/rellenar_plan.py [--fecha AAAA-MM-DD]
Necesita openpyxl.  Lee tres archivos de resultados:
  backend/tests/resultados_plan.json          pytest (parte «srv»; también los casos de caja blanca CB-xx)
  frontend/pruebas/resultados_plan_ui.json    node pruebas/e2e_plan.cjs (parte «ui», en el navegador)
  frontend/pruebas/resultados_plan_cli.json   node --test pruebas/plan_cliente.test.mjs (parte «cli», en el dispositivo)
Reglas: un caso es «Aprobado» solo si pasaron todas las partes que exige; «Falló» si alguna falló;
«Pendiente» si falta alguna parte o el caso necesita una herramienta o una persona (ver PENDIENTES).
"""
import argparse
import json
import os
from datetime import date

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(RAIZ, "docs", "14_Matriz_Trazabilidad_y_Plan_Pruebas_final.xlsx")
ARCHIVOS = {"srv": "backend/tests/resultados_plan.json", "ui": "frontend/pruebas/resultados_plan_ui.json",
            "cli": "frontend/pruebas/resultados_plan_cli.json"}

# Casos que se prueban en más de un lugar: el caso solo se aprueba si pasan todas las partes.
MIXTOS = {
    "CP-RF-01-04": {"srv", "ui"}, "CP-RF-12-01": {"srv", "ui"}, "CP-RF-13-01": {"srv", "cli"}, "CP-RF-13-02": {"srv", "ui"},
    "CP-RF-13-03": {"cli", "ui"}, "CP-RF-14-08": {"srv", "cli"}, "CP-RF-18-01": {"srv", "cli"}, "CP-RF-18-02": {"srv", "cli"},
    "CP-RF-18-03": {"srv", "cli"}, "CP-RF-18-04": {"srv", "cli"}, "CP-RF-23-05": {"srv", "ui"}, "CP-RF-24-02": {"srv", "cli"},
    "CP-RF-25-05": {"srv", "cli", "ui"}, "CP-RF-27-03": {"srv", "ui"}, "CP-RF-28-01": {"srv", "cli"}, "CP-RF-28-02": {"srv", "cli"},
    "CP-RF-28-03": {"srv", "cli"}, "CP-RF-28-08": {"srv", "cli", "ui"}, "CP-RF-37-01": {"srv", "ui"}, "CP-RNF-22": {"srv", "cli"},
    "CP-RNF-25": {"srv", "ui"},
}
# Casos que ninguna prueba automática reemplaza: se dejan «Pendiente» con el motivo y la evidencia parcial que sí existe.
PENDIENTES = {
    "CP-RNF-01": "Medido en el entorno de desarrollo (Chromium sin pantalla, procesador 4 veces más lento): ver e2e_plan.cjs. Falta repetirlo en el equipo de 2 GB de RAM que exige el plan.",
    "CP-RNF-05": "Requiere a 5 agricultores reales haciendo las tareas sin ayuda (cronómetro y observación). No se puede automatizar.",
    "CP-RNF-07": "La medición automática en 7 pantallas pasa (ver e2e_plan.cjs). Faltan Lighthouse (Chrome) y la revisión manual de WCAG 2.1 AA con personas. Quedan etiquetas de navegación, indicadores y marcos de 16 px (el resto del texto es de 18 px).",
    "CP-RNF-09": "Inyección de «<script>», «<img onerror>» y «' OR 1=1 --» probada en formularios y servidor (ver pruebas). Falta ejecutar OWASP ZAP contra el servidor publicado.",
    "CP-RNF-10": "La parte del dispositivo pasa (6 de 6 funciones sin conexión). Falta medir la disponibilidad mensual del servidor publicado (meta 99,0 %).",
    "CP-RNF-14": "Revisión de consultas hecha: todo el acceso a datos usa el ORM (SQLAlchemy); solo hay dos sentencias fijas sin datos del usuario (PRAGMA y SELECT 1) y las cabeceras HSTS y CSP están activas. Falta SSL Labs, que solo funciona contra el sitio publicado con HTTPS.",
    "CP-GEN-01": "Los 28 casos CB-01 a CB-28 pasan y las 4 funciones tienen 100 % de líneas cubiertas (acceso.py y formulas.py). Falta el Quality Gate de SonarQube (duplicación ≤ 5 %), que se ejecuta en CI con SONAR_TOKEN.",
    "CP-GEN-02": "Cobertura medida de forma aproximada: validaciones.py 89 % de ramas y validaciones.js 91 % (node:test); la meta es ≥ 80 %. Falta confirmarlo con pytest-cov --cov-branch, que corre en CI.",
    "CP-GEN-03": "Es del equipo: las bitácoras de tiempo (PSP) y el registro de defectos se llenan durante las iteraciones (documento 23). No se inventan datos de tiempo.",
    "CP-GEN-04": "Se ejecuta antes de la sustentación sobre el sitio publicado (3 min, 6 pasos del Apéndice I del ERS). Los mismos pasos corren hoy en local en e2e.cjs.",
}


def cargar():
    partes = {}
    for k, ruta in ARCHIVOS.items():
        ruta = os.path.join(RAIZ, ruta)
        partes[k] = json.load(open(ruta, encoding="utf-8")) if os.path.exists(ruta) else {}
    return partes


def decidir(caso, partes):
    """Devuelve (estado, texto) para un caso."""
    hallado = {k: p[caso] for k, p in partes.items() if caso in p}
    exigidas = MIXTOS.get(caso, set(hallado) or {"srv"})
    nombres = {"srv": "Servidor", "ui": "Pantalla", "cli": "Dispositivo"}
    if len(hallado) == 1 and caso not in MIXTOS:
        r = next(iter(hallado.values()))
        texto = r["obtenido"]
    else:
        def limpio(t):
            for pref in ("Servidor: ", "Pantalla: ", "Dispositivo: "):
                if t.startswith(pref):
                    return t[len(pref):]
            return t
        texto = " | ".join(f"{nombres[k]}: {limpio(r['obtenido'])}" for k, r in hallado.items())
    if caso in PENDIENTES:
        if any(r["estado"] == "Falló" for r in hallado.values()):
            return "Falló", texto
        return "Pendiente", (texto + " | " if texto else "") + "Pendiente: " + PENDIENTES[caso]
    if any(r["estado"] == "Falló" for r in hallado.values()):
        return "Falló", texto
    faltan = exigidas - set(hallado)
    if faltan:
        return "Pendiente", texto + " | Falta ejecutar: " + ", ".join(nombres[k] for k in sorted(faltan))
    if any(r["estado"] == "Pendiente" for r in hallado.values()):
        return "Pendiente", texto
    return "Aprobado", texto


def rellenar(hoja, col_estado, col_obtenido, partes, contar):
    cuenta = {"Aprobado": 0, "Falló": 0, "Pendiente": 0}
    for fila in range(2, hoja.max_row + 1):
        caso = hoja.cell(fila, 1).value
        if not caso:
            continue
        if caso in PENDIENTES and caso not in {c for p in partes.values() for c in p}:
            estado, texto = "Pendiente", PENDIENTES[caso]
        else:
            estado, texto = decidir(caso, partes)
        hoja.cell(fila, col_estado).value = estado
        hoja.cell(fila, col_obtenido).value = texto if estado != "Pendiente" or texto else None
        cuenta[estado] += 1
    return cuenta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fecha", default=date.today().isoformat())
    a = ap.parse_args()
    partes = cargar()
    libro = openpyxl.load_workbook(EXCEL)
    plan = rellenar(libro["Plan_Pruebas"], 6, 7, partes, True)
    cb = rellenar(libro["Caja_Blanca"], 7, 8, partes, True)
    resumen = libro["Resumen"]
    # indicadores con fórmulas, para que se actualicen si alguien cambia un Estado a mano
    for fila in range(resumen.max_row, 10, -1):
        if resumen.cell(fila, 1).value and str(resumen.cell(fila, 1).value).startswith(("Casos", "Última", "Caja blanca")):
            resumen.cell(fila, 1).value = None
            resumen.cell(fila, 2).value = None
    base = resumen.max_row + 1
    base = max(base, 11)
    filas = [("Casos de caja negra aprobados", '=COUNTIF(Plan_Pruebas!F2:F230,"Aprobado")'),
             ("Casos de caja negra fallidos", '=COUNTIF(Plan_Pruebas!F2:F230,"Falló")'),
             ("Casos de caja negra pendientes", '=COUNTIF(Plan_Pruebas!F2:F230,"Pendiente")'),
             ("Caja blanca aprobados", '=COUNTIF(Caja_Blanca!G2:G29,"Aprobado")'),
             ("Caja blanca fallidos", '=COUNTIF(Caja_Blanca!G2:G29,"Falló")'),
             ("Caja blanca pendientes", '=COUNTIF(Caja_Blanca!G2:G29,"Pendiente")'),
             ("Última ejecución del plan", a.fecha)]
    for i, (k, v) in enumerate(filas):
        resumen.cell(base + i, 1).value = k
        resumen.cell(base + i, 2).value = v
        for c in (1, 2):
            resumen.cell(base + i, c)._style = resumen.cell(10, c)._style
    libro.save(EXCEL)
    print("Caja negra:", plan, "| Caja blanca:", cb)


if __name__ == "__main__":
    main()
